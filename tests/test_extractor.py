"""Tests for PDFHarvest extraction and export."""
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pdfharvest.extractor import extract_pdf
from pdfharvest.exporter import export_to_excel


def test_extract_pdf_not_found():
    """Test that extracting a missing file returns an error gracefully."""
    doc = extract_pdf("non_existent_file.pdf")
    assert doc.error == "File not found"
    assert doc.page_count == 0


def test_extract_and_export(tmp_path):
    """Test real PDF extraction and Excel export."""
    # We use the dummy.pdf we downloaded earlier
    pdf_path = Path("examples/dummy.pdf")
    
    if not pdf_path.exists():
        pytest.skip("dummy.pdf not found in examples/ - skipping test")

    # 1. Extract
    doc = extract_pdf(pdf_path)
    assert doc.error is None
    assert doc.page_count == 1
    assert "Dummy PDF file" in doc.pages[0].text

    # 2. Export to a temporary Excel file
    output_file = tmp_path / "test_output.xlsx"
    export_to_excel([doc], output_file)

    # 3. Check the Excel file
    assert output_file.exists()
    
    wb = load_workbook(output_file)
    assert "dummy_Text" in wb.sheetnames
    
    sheet = wb["dummy_Text"]
    # Row 1 is headers, Row 2 is the first page data
    assert sheet.cell(row=1, column=1).value == "Page"
    assert sheet.cell(row=2, column=1).value == 1
    assert "Dummy PDF file" in sheet.cell(row=2, column=2).value